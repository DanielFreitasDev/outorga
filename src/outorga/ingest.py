"""Ingestão: PDF (texto ou digitalizado), DOCX, imagens e texto puro → `Documento` com páginas.

Páginas sem camada de texto passam por OCR (Tesseract, português; `vendor/tessdata` ou o do sistema), e a confiança
média do OCR vira um sinal no score (documento ilegível não pode cair na faixa verde).
"""

from __future__ import annotations

import hashlib
import os
import re
import subprocess
from dataclasses import dataclass, field
from pathlib import Path

import pymupdf

RAIZ = Path(__file__).resolve().parents[2]
TESSDATA = Path(os.environ.get("OUTORGA_TESSDATA", RAIZ / "vendor" / "tessdata"))
MIN_CHARS_PAGINA = 40


@dataclass
class Documento:
    id: str
    arquivo: str
    nome: str
    paginas: list[str]
    sha256: str
    metodo: str                       # texto | ocr | misto | docx | txt
    ocr_confianca: float | None = None
    offsets: list[tuple[int, int]] = field(default_factory=list)

    @property
    def texto(self) -> str:
        return SEPARADOR.join(self.paginas)

    def pagina_de(self, pos: int) -> int:
        for i, (a, b) in enumerate(self.offsets):
            if a <= pos < b + len(SEPARADOR):
                return i + 1
        return len(self.paginas)


SEPARADOR = "\n\n"


def _calc_offsets(paginas: list[str]) -> list[tuple[int, int]]:
    out, pos = [], 0
    for p in paginas:
        out.append((pos, pos + len(p)))
        pos += len(p) + len(SEPARADOR)
    return out


def limpar(texto: str) -> str:
    texto = texto.replace("­", "").replace("ﬁ", "fi").replace("ﬂ", "fl").replace(" ", " ")
    # hifenização de fim de linha: "constitui-\nção" -> "constituição" (preserva "ATPV-e", "sub-rogação")
    texto = re.sub(r"(\w)-\n([a-zà-ú])", r"\1\2", texto)
    texto = re.sub(r"[ \t]+\n", "\n", texto)
    return re.sub(r"\n{3,}", "\n\n", texto).strip()


def ocr_imagem(png: bytes) -> tuple[str, float]:
    """OCR com Tesseract (por). Retorna (texto, confiança média 0-100)."""
    env = dict(os.environ, OMP_THREAD_LIMIT="4")
    if (TESSDATA / "por.traineddata").exists():  # dados próprios; senão, os do sistema (pacote tesseract-ocr-por)
        env["TESSDATA_PREFIX"] = str(TESSDATA)
    r = subprocess.run(["tesseract", "stdin", "stdout", "-l", "por", "--psm", "4", "-c", "tessedit_create_tsv=1"], input=png,
                       capture_output=True, env=env, timeout=300)
    if r.returncode != 0:
        raise RuntimeError(f"tesseract falhou: {r.stderr.decode(errors='replace')[:300]}")
    linhas: dict[tuple, list[str]] = {}
    confs = []
    for ln in r.stdout.decode("utf-8", errors="replace").splitlines()[1:]:
        cols = ln.split("\t")
        if len(cols) < 12 or cols[0] != "5":
            continue
        chave = (int(cols[2]), int(cols[3]), int(cols[4]))
        palavra = cols[11]
        if palavra.strip():
            linhas.setdefault(chave, []).append(palavra)
            try:
                c = float(cols[10])
                if c >= 0:
                    confs.append(c)
            except ValueError:
                pass
    texto, ultimo_bloco = [], None
    for (bloco, par, _), palavras in linhas.items():
        if ultimo_bloco is not None and (bloco, par) != ultimo_bloco:
            texto.append("")
        texto.append(" ".join(palavras))
        ultimo_bloco = (bloco, par)
    return "\n".join(texto), (sum(confs) / len(confs) if confs else 0.0)


def _ler_pdf(caminho: Path) -> tuple[list[str], str, float | None]:
    doc = pymupdf.open(caminho)
    paginas, metodos, confs = [], set(), []
    for page in doc:
        t = page.get_text("text")
        if len(t.strip()) >= MIN_CHARS_PAGINA:
            paginas.append(limpar(t))
            metodos.add("texto")
        else:
            png = page.get_pixmap(dpi=300, colorspace=pymupdf.csGRAY).tobytes("png")
            t, c = ocr_imagem(png)
            paginas.append(limpar(t))
            confs.append(c)
            metodos.add("ocr")
    metodo = "misto" if len(metodos) > 1 else (metodos.pop() if metodos else "texto")
    return paginas, metodo, (sum(confs) / len(confs) if confs else None)


def _ler_docx(caminho: Path) -> list[str]:
    import docx

    d = docx.Document(str(caminho))
    partes = [p.text for p in d.paragraphs]
    for tabela in d.tables:
        for linha in tabela.rows:
            partes.append(" | ".join(c.text.strip() for c in linha.cells))
    return [limpar("\n".join(partes))]


def carregar(caminho: str | Path, doc_id: str | None = None) -> Documento:
    caminho = Path(caminho)
    dados = caminho.read_bytes()
    sha = hashlib.sha256(dados).hexdigest()
    suf = caminho.suffix.lower()
    conf = None
    if suf == ".pdf":
        paginas, metodo, conf = _ler_pdf(caminho)
    elif suf == ".docx":
        paginas, metodo = _ler_docx(caminho), "docx"
    elif suf in (".png", ".jpg", ".jpeg", ".tif", ".tiff", ".webp", ".bmp"):
        from PIL import Image
        import io

        img = Image.open(io.BytesIO(dados)).convert("L")
        buf = io.BytesIO()
        img.save(buf, format="PNG")
        t, conf = ocr_imagem(buf.getvalue())
        paginas, metodo = [limpar(t)], "ocr"
    else:
        paginas, metodo = [limpar(dados.decode("utf-8", errors="replace"))], "txt"
    d = Documento(doc_id or caminho.stem, str(caminho), caminho.name, paginas, sha, metodo, conf)
    d.offsets = _calc_offsets(d.paginas)
    return d
